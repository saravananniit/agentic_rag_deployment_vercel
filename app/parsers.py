"""File -> plain text, and text -> chunks. Light dependencies only (Vercel bundle limit)."""
from __future__ import annotations

import csv
import io
import json
import re

SUPPORTED = {
    ".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md", ".html", ".htm", ".json", ".log",
}

# NV-Embed-QA accepts ~512 tokens; ~1000 characters stays safely below that.
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150


class UnsupportedFile(ValueError):
    pass


def extension(filename: str) -> str:
    i = filename.rfind(".")
    return filename[i:].lower() if i >= 0 else ""


def extract_text(filename: str, data: bytes) -> str:
    ext = extension(filename)
    if ext not in SUPPORTED:
        raise UnsupportedFile(
            f"Unsupported file type '{ext or 'unknown'}'. Supported: {', '.join(sorted(SUPPORTED))}"
        )
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        out = []
        for n, slide in enumerate(prs.slides, 1):
            out.append(f"Slide {n}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    out.append(shape.text_frame.text)
        return "\n".join(out)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    out.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(out)
    if ext == ".csv":
        text = _decode(data)
        rows = csv.reader(io.StringIO(text))
        return "\n".join(" | ".join(r) for r in rows)
    if ext in (".html", ".htm"):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(_decode(data), "html.parser")
        for t in soup(["script", "style", "noscript"]):
            t.decompose()
        return soup.get_text("\n")
    if ext == ".json":
        try:
            return json.dumps(json.loads(_decode(data)), indent=2)
        except json.JSONDecodeError:
            return _decode(data)
    return _decode(data)


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Paragraph-aware chunking with a small character overlap."""
    text = re.sub(r"[ \t]+", " ", text).replace("\r", "")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for p in paragraphs:
        # break oversized paragraphs on sentence/space boundaries
        while len(p) > size:
            cut = max(p.rfind(". ", 0, size), p.rfind("\n", 0, size), p.rfind(" ", 0, size))
            cut = cut + 1 if cut > size // 2 else size
            piece, p = p[:cut].strip(), p[cut:].strip()
            if current:
                chunks.append(current)
                current = ""
            chunks.append(piece)
        if not p:
            continue
        if len(current) + len(p) + 2 <= size:
            current = f"{current}\n\n{p}" if current else p
        else:
            if current:
                chunks.append(current)
            current = p
    if current:
        chunks.append(current)
    # add overlap from the tail of the previous chunk
    if overlap and len(chunks) > 1:
        merged = [chunks[0]]
        for prev, cur in zip(chunks, chunks[1:]):
            tail = prev[-overlap:]
            sp = tail.find(" ")
            tail = tail[sp + 1:] if sp != -1 else tail
            merged.append(f"{tail} {cur}"[: size + overlap])
        chunks = merged
    return chunks
